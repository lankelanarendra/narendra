import openpyxl

file="C:\\Users\\user\\PycharmProjects\\python_tutorials\\pythonProject1\\pythonProject21\\student_data.xlsx"

workbook=openpyxl.load_workbook(file)
sheet=workbook["Sheet1"]

rows=sheet.max_row  #count number of rows in a excel sheet
cols=sheet.max_column #count number of column in a excel sheet

# reading all the rows and column from the excel

for i in range(1,rows+1):
    for j in range(1,cols+1):
        print(sheet.cell(i,j).value)
    print()





