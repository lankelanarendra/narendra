import openpyxl

file="C:\\Users\\user\\Documents\\Book1.xlsx"

workbook=openpyxl.load_workbook(file)
sheet=workbook.active

for i in range(1,6):
    for j in range(1,4):
        sheet.cell(i,j).value="selenium"
workbook.save(file)
#
# multiple data into a excel sheet

# import openpyxl
#
# file="C:\\Users\\user\\Documents\\testcase.xlsx"
#
# workbook=openpyxl.load_workbook(file)
# sheet=workbook.active
# sheet.cell(1,1).value=123
# sheet.cell(1,2).value="sai"
# sheet.cell(1,3).value="Tractor"
#
# sheet.cell(2,1).value=234
# sheet.cell(2,2).value="pavan"
# sheet.cell(2,3).value="milk auto"
#
# sheet.cell(3,1).value=321
# sheet.cell(3,2).value="raghu"
# sheet.cell(3,3).value="college"
#
# workbook.save(file)




